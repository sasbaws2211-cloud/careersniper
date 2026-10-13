"""Turn an uploaded résumé file (PDF, Word .docx or plain text) into text.

Only the text is kept: the file itself is never stored. Parsing runs in a worker thread (it is CPU work on untrusted input) and is
bounded: file size, PDF page count and the unpacked size of a .docx are all capped.
"""
from __future__ import annotations

import asyncio
import io
import zipfile

from .resume import ResumeError

MAX_FILE_BYTES = 2_000_000          # the file; the route allows a little more for the multipart envelope
MAX_PDF_PAGES = 15
MAX_UNPACKED = 20_000_000           # a .docx is a zip: refuse one that unpacks to something huge
MIN_TEXT = 50                       # less than this means the file is an image, not text

ACCEPT = ".pdf,.docx,.txt"
UNREADABLE = "We could not read that file. Try another copy, or paste the text of your résumé below."


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                raise ResumeError("That PDF is password-protected. Remove the password and upload it again, or paste the text instead.")
        except ResumeError:
            raise
        except Exception:
            raise ResumeError("That PDF is password-protected. Remove the password and upload it again, or paste the text instead.")
    return "\n".join((page.extract_text() or "") for page in reader.pages[:MAX_PDF_PAGES])


def _docx(data: bytes) -> str:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise ResumeError(UNREADABLE)
    if sum(i.file_size for i in zf.infolist()) > MAX_UNPACKED or "word/document.xml" not in zf.namelist():
        raise ResumeError(UNREADABLE)
    import docx
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:                       # many résumé templates put sections in tables
        for row in table.rows:
            cells = []
            for c in row.cells:
                t = " ".join(c.text.split())
                if t and t not in cells:
                    cells.append(t)
            if cells:
                parts.append("  |  ".join(cells))
    return "\n".join(parts)


def _extract(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if not data:
        raise ResumeError("That file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ResumeError(f"That file is too large. The limit is {MAX_FILE_BYTES // 1_000_000} MB; a résumé is normally well under 1 MB.")
    try:
        if data[:5] == b"%PDF-":
            text = _pdf(data)
        elif data[:2] == b"PK" and name.endswith(".docx"):
            text = _docx(data)
        elif name.endswith((".doc", ".rtf", ".pages", ".odt")):
            raise ResumeError("That file type is not supported. Save it as a PDF or a Word .docx file (or paste the text) and try again.")
        elif name.endswith((".txt", ".md", ".text")):
            if b"\x00" in data[:2000]:
                raise ResumeError(UNREADABLE)                   # a "text" file that is really binary
            text = data.decode("utf-8", errors="replace")
        else:
            raise ResumeError("Upload a PDF, a Word .docx file or a plain text file.")
    except ResumeError:
        raise
    except Exception:
        raise ResumeError(UNREADABLE)
    if len(text.strip()) < MIN_TEXT:
        raise ResumeError("We could not find any text in that file. It may be a scan or an image. Upload a text-based PDF or Word file, "
                          "or paste the text of your résumé below.")
    return text


async def extract_text(filename: str, data: bytes) -> str:
    """The résumé's text, or ResumeError with a message the member can act on."""
    return await asyncio.to_thread(_extract, filename, data)
